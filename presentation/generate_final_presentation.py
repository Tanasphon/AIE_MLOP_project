from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


WORKSPACE_DIR = Path(__file__).resolve().parents[2]
PROJECT_DIR = WORKSPACE_DIR / "project"
OUTPUT_PATH = WORKSPACE_DIR / "final_mobilevit_mlops_presentation.pptx"
PROJECT_OUTPUT_PATH = PROJECT_DIR / "presentation" / "mobilevit_mlops_presentation.pptx"
SAMPLE_IMAGE = PROJECT_DIR / "sample_images" / "benchmark.png"

FONT_NAME = "Tahoma"
NAVY = RGBColor(22, 54, 92)
BLUE = RGBColor(50, 117, 181)
LIGHT_BLUE = RGBColor(221, 235, 247)
DARK = RGBColor(40, 40, 40)
GREY = RGBColor(89, 89, 89)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(112, 173, 71)
ORANGE = RGBColor(237, 125, 49)
PURPLE = RGBColor(112, 48, 160)


def set_run(run, size=18, bold=False, color=DARK):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(12.0), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    set_run(p.runs[0], size=28, bold=True, color=NAVY)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.12), Inches(12.0), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.68), Inches(1.22), Inches(11.8), Inches(0.35))
        p2 = sub.text_frame.paragraphs[0]
        p2.text = subtitle
        set_run(p2.runs[0], size=13, color=GREY)


def add_footer(slide, page: int):
    box = slide.shapes.add_textbox(Inches(10.55), Inches(7.05), Inches(2.1), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = f"MobileViT MLOps | {page}"
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.runs[0], size=9, color=GREY)


def add_bullets(slide, items: list[str], x=0.85, y=1.55, w=11.8, h=4.7, size=19):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(8)
        p.font.name = FONT_NAME
        p.font.size = Pt(size)
        p.font.color.rgb = DARK


def add_card(slide, x, y, w, h, title, body, color=LIGHT_BLUE, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = RGBColor(180, 180, 180)

    box = slide.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.12), Inches(w - 0.32), Inches(h - 0.2))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    set_run(p.runs[0], size=title_size, bold=True, color=NAVY)
    p2 = tf.add_paragraph()
    p2.text = body
    p2.space_before = Pt(5)
    p2.font.name = FONT_NAME
    p2.font.size = Pt(body_size)
    p2.font.color.rgb = DARK


def add_table(slide, rows, x, y, w, h, font_size=12):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for col_idx in range(len(rows[0])):
        table.columns[col_idx].width = Inches(w / len(rows[0]))

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = BLUE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    set_run(run, size=font_size, bold=(r_idx == 0), color=WHITE if r_idx == 0 else DARK)
    return table


def add_step(slide, number, title, detail, x, y, color=BLUE):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.55), Inches(0.55))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    num_box = slide.shapes.add_textbox(Inches(x), Inches(y + 0.08), Inches(0.55), Inches(0.3))
    p = num_box.text_frame.paragraphs[0]
    p.text = str(number)
    p.alignment = PP_ALIGN.CENTER
    set_run(p.runs[0], size=15, bold=True, color=WHITE)
    add_card(slide, x + 0.7, y - 0.05, 3.25, 0.95, title, detail, RGBColor(242, 247, 252), 13, 11)


def add_arrow(slide, x, y, w=0.65, h=0.25, color=GREY):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = color
    arrow.line.fill.background()


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Cover
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BLUE
    bg.line.fill.background()
    title = slide.shapes.add_textbox(Inches(0.75), Inches(1.05), Inches(11.9), Inches(1.15))
    p = title.text_frame.paragraphs[0]
    p.text = "High-Throughput Image Classification Service"
    p.alignment = PP_ALIGN.CENTER
    set_run(p.runs[0], size=34, bold=True, color=NAVY)
    sub = slide.shapes.add_textbox(Inches(1.05), Inches(2.38), Inches(11.2), Inches(0.95))
    p2 = sub.text_frame.paragraphs[0]
    p2.text = "ระบบ API จำแนกรูปภาพด้วย FastAPI และโมเดล apple/mobilevit-small"
    p2.alignment = PP_ALIGN.CENTER
    set_run(p2.runs[0], size=20, color=GREY)
    add_card(slide, 2.0, 4.1, 9.4, 0.9, "สมาชิกกลุ่ม", "ธนัสถ์ภณ อ่างทอง | จีฮาน สุทธินิพนธ์นาม | อกัณห์ เกษเพชร", WHITE, 16, 15)
    add_footer(slide, 1)

    # 2 Objective
    slide = prs.slides.add_slide(blank)
    add_title(slide, "วัตถุประสงค์ของโปรเจกต์")
    add_bullets(slide, [
        "พัฒนา API สำหรับจำแนกรูปภาพที่สามารถเรียกใช้งานได้จริง",
        "ปรับแต่งโมเดลให้เล็กลงและเหมาะกับ CPU inference",
        "รองรับการเรียกใช้งานพร้อมกันโดยไม่ทำให้ API ค้าง",
        "เตรียม Docker, CI/CD, Postman, JMeter, รายงาน และสไลด์นำเสนอให้ครบตามโจทย์",
    ], size=21)
    add_footer(slide, 2)

    # 3 Steps timeline
    slide = prs.slides.add_slide(blank)
    add_title(slide, "ขั้นตอนการทำโปรเจกต์", "ลำดับงานตั้งแต่เลือกโมเดลจนถึงเตรียมส่งงาน")
    add_step(slide, 1, "เลือกโมเดล", "ใช้ apple/mobilevit-small จาก Hugging Face", 0.75, 1.75, BLUE)
    add_step(slide, 2, "Baseline", "ทดสอบ PyTorch เพื่อเก็บขนาดและ latency", 4.55, 1.75, GREEN)
    add_step(slide, 3, "ONNX", "แปลงโมเดลให้เหมาะกับ inference", 8.35, 1.75, ORANGE)
    add_arrow(slide, 4.15, 2.05)
    add_arrow(slide, 7.95, 2.05)
    add_step(slide, 4, "Quantization", "ลดขนาดโมเดลด้วย Dynamic Quantization", 0.75, 4.1, PURPLE)
    add_step(slide, 5, "FastAPI + Docker", "สร้าง API, validation และ container", 4.55, 4.1, BLUE)
    add_step(slide, 6, "Test + Report", "ทดสอบ API, load test และสรุปผล", 8.35, 4.1, GREEN)
    add_arrow(slide, 4.15, 4.4)
    add_arrow(slide, 7.95, 4.4)
    add_footer(slide, 3)

    # 4 Workflow visual
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Workflow การทำงานของระบบ")
    labels = [
        ("รูปภาพ", "ผู้ใช้ส่งไฟล์ภาพจริง"),
        ("Validation", "ตรวจชนิดไฟล์ ขนาด และไฟล์เสีย"),
        ("Inference", "ประมวลผลด้วย ONNX Quantized"),
        ("Response", "ส่ง label และ score เป็น JSON"),
    ]
    xs = [0.8, 3.9, 7.0, 10.1]
    for i, ((t, b), x) in enumerate(zip(labels, xs)):
        add_card(slide, x, 2.05, 2.35, 1.35, t, b, RGBColor(232, 242, 250), 15, 12)
        if i < len(labels) - 1:
            add_arrow(slide, x + 2.42, 2.6, 0.72, 0.28, BLUE)
    add_bullets(slide, [
        "FastAPI ทำหน้าที่รับ request และควบคุม response",
        "ProcessPoolExecutor แยกงาน inference ที่ใช้ CPU ออกจาก event loop",
        "Docker ทำให้ระบบรันซ้ำได้ง่ายทั้ง local และ cloud",
    ], y=4.35, size=18)
    add_footer(slide, 4)

    # 5 Model
    slide = prs.slides.add_slide(blank)
    add_title(slide, "โมเดลที่เลือกใช้", "apple/mobilevit-small จาก Hugging Face")
    add_card(slide, 0.85, 1.7, 3.6, 1.35, "Lightweight", "ขนาดไม่ใหญ่มาก เหมาะกับ CPU inference")
    add_card(slide, 4.85, 1.7, 3.6, 1.35, "CNN + Transformer", "เรียนรู้ทั้งรายละเอียดเฉพาะจุดและภาพรวม")
    add_card(slide, 8.85, 1.7, 3.6, 1.35, "API Ready", "นำไปใช้กับระบบ FastAPI และ Docker ได้")
    add_bullets(slide, [
        "งานที่ใช้: Image Classification",
        "Runtime ที่ใช้จริง: ONNX Quantized",
        "รูปภาพทดสอบมาจาก Hugging Face documentation-images",
    ], y=3.75, size=19)
    add_footer(slide, 5)

    # 6 Architecture
    slide = prs.slides.add_slide(blank)
    add_title(slide, "System Architecture")
    parts = [
        ("Client", "curl / Postman / JMeter", 0.55),
        ("FastAPI", "รับ request และตรวจสอบ input", 3.05),
        ("ProcessPoolExecutor", "แยกงาน inference", 5.7),
        ("ONNX Quantized", "ประมวลผลโมเดล", 8.5),
        ("JSON Response", "ส่งผลลัพธ์กลับ", 10.95),
    ]
    for i, (t, b, x) in enumerate(parts):
        add_card(slide, x, 2.1, 2.0, 1.25, t, b, RGBColor(232, 242, 250), 14, 11)
        if i < len(parts) - 1:
            add_arrow(slide, x + 2.05, 2.58, 0.45, 0.23, BLUE)
    add_bullets(slide, [
        "API เป็น async request layer เพื่อรองรับการรับส่งข้อมูล",
        "Inference เป็นงาน CPU-bound จึงแยกไปทำใน process pool",
        "โมเดลและแอปถูกบรรจุใน Docker เพื่อให้ deploy ได้ง่าย",
    ], y=4.2, size=18)
    add_footer(slide, 6)

    # 7 Optimization
    slide = prs.slides.add_slide(blank)
    add_title(slide, "ผลการ Optimization")
    add_table(slide, [
        ["Runtime", "Model Size", "Avg Latency", "P95 Latency"],
        ["PyTorch", "21.44 MB", "34.43 ms", "37.51 ms"],
        ["ONNX", "21.49 MB", "55.58 ms", "63.64 ms"],
        ["ONNX Quantized", "11.56 MB", "47.00 ms", "56.47 ms"],
    ], 0.9, 1.55, 11.5, 1.65)
    add_card(slide, 1.0, 4.05, 5.2, 1.35, "ผลลัพธ์สำคัญ", "Quantized ลดขนาดโมเดลลงประมาณ 46% จาก ONNX ปกติ", GREEN)
    add_card(slide, 7.0, 4.05, 5.2, 1.35, "ข้อสังเกต", "Latency ของ ONNX Quantized ดีกว่า ONNX ปกติในเครื่องทดสอบ", ORANGE)
    add_footer(slide, 7)

    # 8 Result source
    slide = prs.slides.add_slide(blank)
    add_title(slide, "ที่มาของข้อมูลและผลลัพธ์", "อธิบายว่าตัวเลขในสไลด์มาจากการทดสอบส่วนใด")
    add_table(slide, [
        ["ผลลัพธ์", "ที่มา", "นำไปใช้ในสไลด์"],
        ["Model Size", "ขนาดไฟล์โมเดลในโฟลเดอร์ models", "เปรียบเทียบ PyTorch / ONNX / Quantized"],
        ["Latency", "จับเวลาการ inference ด้วยรูปภาพ benchmark.png", "ประเมินความเร็วของแต่ละ runtime"],
        ["Prediction", "เรียก endpoint /predict ด้วยรูปภาพจริง", "ยืนยันว่า API ใช้งานได้จริง"],
        ["Load Test", "ส่ง request พร้อมกันหลายรายการ", "วิเคราะห์ throughput และ P95 latency"],
    ], 0.8, 1.6, 11.7, 2.7, font_size=10)
    add_bullets(slide, [
        "รูปภาพทดสอบเป็นภาพแมวจาก Hugging Face documentation-images",
        "ผลลัพธ์ benchmark ถูกบันทึกไว้ใน models/benchmark_results.json",
        "ผล API ยืนยันจาก JSON response ที่มี model: apple/mobilevit-small และ runtime: onnx-quantized",
    ], y=4.75, size=16)
    add_footer(slide, 8)

    # 9 Model measurement methodology
    slide = prs.slides.add_slide(blank)
    add_title(slide, "วิธีวัดประสิทธิภาพของโมเดล")
    add_card(slide, 0.8, 1.55, 3.8, 1.35, "1. Baseline", "โหลดโมเดล PyTorch ต้นฉบับและวัดขนาด/เวลา inference", RGBColor(232, 242, 250))
    add_card(slide, 4.8, 1.55, 3.8, 1.35, "2. ONNX", "แปลงโมเดลเป็น ONNX แล้ววัด latency ด้วย ONNX Runtime", RGBColor(232, 242, 250))
    add_card(slide, 8.8, 1.55, 3.8, 1.35, "3. Quantized", "ทำ Dynamic Quantization และวัดผลซ้ำด้วย input เดียวกัน", RGBColor(232, 242, 250))
    add_table(slide, [
        ["Metric", "ความหมาย", "เหตุผลที่ใช้วัด"],
        ["Model Size", "ขนาดไฟล์โมเดล", "ดูว่า optimization ลดขนาดได้เท่าไร"],
        ["Average Latency", "เวลา inference เฉลี่ย", "ดูความเร็วโดยรวม"],
        ["P95 Latency", "เวลาที่ 95% ของ request ไม่เกินค่านี้", "ดูความเสถียรของ response time"],
    ], 1.0, 3.75, 11.3, 1.65, font_size=10)
    add_footer(slide, 9)

    # 10 API/load measurement methodology
    slide = prs.slides.add_slide(blank)
    add_title(slide, "วิธีวัดประสิทธิภาพของ API")
    add_bullets(slide, [
        "ทดสอบ API ด้วย Docker เพื่อให้ใกล้เคียงสภาพแวดล้อม deploy จริง",
        "ส่งไฟล์รูปภาพจริงเข้า endpoint /predict และตรวจว่า response เป็น JSON ที่ถูกต้อง",
        "ทดสอบ error handling ด้วยไฟล์ที่ไม่ใช่รูปภาพ เพื่อดูว่า API ตอบ 400 Bad Request",
        "ทดสอบโหลดด้วย concurrent requests เพื่อดู throughput, average latency และ P95 latency",
        "ใช้ Error % เป็นตัวบอกว่าระบบเริ่มรับโหลดไม่ไหวหรือไม่",
    ], size=18)
    add_footer(slide, 10)

    # 11 API and validation
    slide = prs.slides.add_slide(blank)
    add_title(slide, "API และ Error Handling")
    add_bullets(slide, [
        "Endpoint หลัก: /predict รับไฟล์รูปภาพแบบ multipart/form-data",
        "รองรับไฟล์ JPEG, PNG และ WebP",
        "ตรวจสอบไฟล์ว่าง ไฟล์เสีย และขนาดไฟล์ก่อนเข้าโมเดล",
        "ตอบกลับ 400 Bad Request เมื่อ input ไม่ถูกต้อง",
        "ตอบกลับ 413 Request Entity Too Large เมื่อไฟล์ใหญ่เกินกำหนด",
    ], size=19)
    add_footer(slide, 11)

    # 12 Testing
    slide = prs.slides.add_slide(blank)
    add_title(slide, "ผลการทดสอบ API")
    add_card(slide, 0.9, 1.55, 3.5, 1.3, "Unit Test", "ผ่าน 3/3 tests", GREEN)
    add_card(slide, 4.9, 1.55, 3.5, 1.3, "Docker Test", "API รันและทำนายผลได้จริง", GREEN)
    add_card(slide, 8.9, 1.55, 3.5, 1.3, "Prediction", "runtime: onnx-quantized", GREEN)
    add_table(slide, [
        ["Input", "ผลลัพธ์"],
        ["sample_images/benchmark.png", "tabby cat, Egyptian cat, tiger cat"],
        ["ไฟล์ text/plain", "400 Bad Request"],
    ], 1.2, 3.6, 10.8, 1.25)
    add_footer(slide, 12)

    # 13 Load test
    slide = prs.slides.add_slide(blank)
    add_title(slide, "ผลการทดสอบโหลด")
    add_table(slide, [
        ["หัวข้อ", "ผลการทดสอบ"],
        ["จำนวน request", "30"],
        ["Concurrent users", "10"],
        ["Success", "30/30"],
        ["Throughput", "10.32 requests/second"],
        ["Average Latency", "874.90 ms"],
        ["P95 Latency", "1187.96 ms"],
    ], 1.3, 1.4, 10.5, 3.0)
    add_bullets(slide, [
        "ระบบตอบสนองครบทุก request ในชุดทดสอบขนาดเล็ก",
        "เมื่อเพิ่ม threads มากขึ้น latency จะสูงขึ้นตามภาระ CPU",
    ], y=5.0, size=17)
    add_footer(slide, 13)

    # 14 CI/CD
    slide = prs.slides.add_slide(blank)
    add_title(slide, "CI/CD Pipeline")
    stages = [("Push Code", "GitHub"), ("Install", "Dependencies"), ("Test", "pytest"), ("Deploy", "Hugging Face Spaces")]
    for i, (t, b) in enumerate(stages):
        x = 1.0 + i * 2.8
        add_card(slide, x, 1.85, 2.25, 1.2, t, b, RGBColor(232, 242, 250))
        if i < 3:
            add_arrow(slide, x + 2.32, 2.35, 0.45, 0.23, BLUE)
    add_bullets(slide, [
        "Workflow อยู่ที่ .github/workflows/ci-cd.yml",
        "รัน unit test ทุกครั้งที่ push หรือ pull request ไปยัง main",
        "รองรับ auto-deploy เมื่อกำหนด HF_TOKEN และ HF_SPACE_REPO_ID",
    ], y=4.15, size=18)
    add_footer(slide, 14)

    # 15 Live demo
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Live Demo: API บน Docker")
    if SAMPLE_IMAGE.exists():
        slide.shapes.add_picture(str(SAMPLE_IMAGE), Inches(0.9), Inches(1.55), width=Inches(4.3))
        cap = slide.shapes.add_textbox(Inches(0.9), Inches(5.35), Inches(4.3), Inches(0.4))
        p = cap.text_frame.paragraphs[0]
        p.text = "ภาพตัวอย่างจริงที่ใช้ทดสอบ endpoint /predict"
        p.alignment = PP_ALIGN.CENTER
        set_run(p.runs[0], size=12, color=GREY)
    add_card(slide, 5.65, 1.55, 6.6, 1.2, "ขั้นตอน Demo", "เปิด Docker API → ส่งรูปภาพ → ตรวจ JSON response", LIGHT_BLUE, 17, 15)
    add_card(slide, 5.65, 3.15, 6.6, 1.2, "ผลลัพธ์ที่ได้", "model: apple/mobilevit-small | runtime: onnx-quantized", GREEN, 17, 14)
    add_card(slide, 5.65, 4.75, 6.6, 1.2, "Top Predictions", "tabby cat, Egyptian cat, tiger cat", ORANGE, 17, 14)
    add_footer(slide, 15)

    # 16 Deliverables
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Deliverables ที่จัดทำ")
    add_bullets(slide, [
        "Source Code Repository บน GitHub",
        "README.md ภาษาไทย พร้อมวิธีใช้งานและ cURL",
        "Dockerfile และโมเดล ONNX Quantized",
        "GitHub Actions สำหรับ CI/CD",
        "Postman Collection และ JMeter Test Plan",
        "รายงาน PDF และสไลด์นำเสนอ",
    ], size=20)
    add_footer(slide, 16)

    # 17 Summary
    slide = prs.slides.add_slide(blank)
    add_title(slide, "สรุปผล")
    add_bullets(slide, [
        "พัฒนา API จำแนกรูปภาพด้วย apple/mobilevit-small ได้สำเร็จ",
        "ระบบใช้งานจริงผ่าน Docker และตอบผลลัพธ์ด้วย ONNX Quantized Runtime",
        "มีการทดสอบ API, error handling, unit test และ load test",
        "เตรียมเอกสารและ pipeline ครบตามข้อกำหนดของงาน",
    ], size=21)
    thanks = slide.shapes.add_textbox(Inches(0.85), Inches(5.75), Inches(11.6), Inches(0.7))
    p = thanks.text_frame.paragraphs[0]
    p.text = "ขอบคุณครับ / ค่ะ"
    p.alignment = PP_ALIGN.CENTER
    set_run(p.runs[0], size=28, bold=True, color=NAVY)
    add_footer(slide, 17)

    prs.save(OUTPUT_PATH)
    prs.save(PROJECT_OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")
    print(f"Created {PROJECT_OUTPUT_PATH}")


if __name__ == "__main__":
    build()
