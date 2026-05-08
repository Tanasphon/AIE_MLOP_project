from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT_DIR = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT_DIR / "presentation" / "mobilevit_mlops_presentation.pptx"

FONT_NAME = "Tahoma"
NAVY = RGBColor(22, 54, 92)
BLUE = RGBColor(50, 117, 181)
LIGHT_BLUE = RGBColor(221, 235, 247)
GREY = RGBColor(89, 89, 89)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(112, 173, 71)
ORANGE = RGBColor(237, 125, 49)


def set_run(run, size=24, bold=False, color=NAVY):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(12.0), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    set_run(p.runs[0], size=28, bold=True, color=NAVY)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.12), Inches(12.0), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.68), Inches(1.22), Inches(11.8), Inches(0.35))
        p2 = sub.text_frame.paragraphs[0]
        p2.text = subtitle
        set_run(p2.runs[0], size=13, color=GREY)


def add_bullets(slide, items: list[str], x=0.85, y=1.55, w=11.8, h=4.7, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(8)
        p.font.name = FONT_NAME
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(40, 40, 40)


def add_card(slide, x, y, w, h, title, body, color=LIGHT_BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = RGBColor(180, 180, 180)

    box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.13), Inches(w - 0.36), Inches(h - 0.24))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    set_run(p.runs[0], size=17, bold=True, color=NAVY)
    p2 = tf.add_paragraph()
    p2.text = body
    p2.space_before = Pt(6)
    p2.font.name = FONT_NAME
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(40, 40, 40)


def add_table(slide, rows, x, y, w, h, font_size=13):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for col_idx in range(len(rows[0])):
        table.columns[col_idx].width = Inches(w / len(rows[0]))

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = BLUE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    set_run(run, size=font_size, bold=(r_idx == 0), color=WHITE if r_idx == 0 else RGBColor(40, 40, 40))
    return table


def add_footer(slide, page: int):
    box = slide.shapes.add_textbox(Inches(10.6), Inches(7.05), Inches(2.0), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = f"MobileViT MLOps | {page}"
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.runs[0], size=9, color=GREY)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BLUE
    bg.line.fill.background()
    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(11.8), Inches(1.2))
    p = title.text_frame.paragraphs[0]
    p.text = "High-Throughput Image Classification Service"
    p.alignment = PP_ALIGN.CENTER
    set_run(p.runs[0], size=34, bold=True, color=NAVY)
    sub = slide.shapes.add_textbox(Inches(1.0), Inches(2.55), Inches(11.3), Inches(1.0))
    p2 = sub.text_frame.paragraphs[0]
    p2.text = "ระบบ API จำแนกรูปภาพด้วย FastAPI และ apple/mobilevit-small"
    p2.alignment = PP_ALIGN.CENTER
    set_run(p2.runs[0], size=20, color=GREY)
    team = slide.shapes.add_textbox(Inches(2.0), Inches(4.15), Inches(9.4), Inches(1.2))
    tf = team.text_frame
    tf.text = "สมาชิก: ธนัสถ์ภณ อ่างทอง | จีฮาน สุทธินิพนธ์นาม | อกัณห์ เกษเพชร"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_run(tf.paragraphs[0].runs[0], size=15, color=NAVY)
    add_footer(slide, 1)

    # 2
    slide = prs.slides.add_slide(blank)
    add_title(slide, "วัตถุประสงค์ของโปรเจกต์")
    add_bullets(
        slide,
        [
            "พัฒนา API สำหรับจำแนกรูปภาพที่สามารถเรียกใช้งานได้จริง",
            "ปรับแต่งโมเดลให้เหมาะกับงาน inference บน CPU ด้วย ONNX และ Quantization",
            "จัดการ request พร้อมกันโดยไม่ทำให้ API ค้าง",
            "เตรียม Docker, CI/CD, Postman และ JMeter สำหรับกระบวนการ MLOps",
        ],
        size=21,
    )
    add_footer(slide, 2)

    # 3
    slide = prs.slides.add_slide(blank)
    add_title(slide, "โมเดลที่เลือกใช้", "apple/mobilevit-small จาก Hugging Face")
    add_card(slide, 0.8, 1.6, 3.8, 1.5, "Lightweight", "ขนาดไม่ใหญ่มาก เหมาะกับ CPU inference")
    add_card(slide, 4.8, 1.6, 3.8, 1.5, "CNN + Transformer", "เรียนรู้ทั้งรายละเอียดเฉพาะจุดและภาพรวม")
    add_card(slide, 8.8, 1.6, 3.8, 1.5, "Real-world API", "เหมาะกับระบบที่ต้องการ response เร็วและ deploy ง่าย")
    add_bullets(
        slide,
        [
            "งานที่ใช้: Image Classification",
            "Runtime ที่ใช้จริงใน API: ONNX Quantized",
            "ผลลัพธ์ที่ API ส่งกลับ: label และ confidence score ในรูปแบบ JSON",
        ],
        y=3.75,
        size=19,
    )
    add_footer(slide, 3)

    # 4
    slide = prs.slides.add_slide(blank)
    add_title(slide, "System Architecture")
    cards = [
        ("Client", "curl / Postman / JMeter", 0.6),
        ("FastAPI", "รับ request และตรวจสอบ input", 3.1),
        ("ProcessPoolExecutor", "แยกงาน inference ที่ใช้ CPU", 5.9),
        ("ONNX Quantized", "ประมวลผลโมเดล", 8.9),
        ("JSON Response", "ส่งผลลัพธ์กลับ", 11.0),
    ]
    for title_text, body, x in cards:
        add_card(slide, x, 2.15, 2.0, 1.25, title_text, body, color=RGBColor(232, 242, 250))
    add_bullets(
        slide,
        [
            "FastAPI ใช้ async request layer เพื่อรองรับการรับส่งข้อมูล",
            "Inference เป็นงาน CPU-bound จึงแยกไปทำงานใน process pool",
            "Docker ใช้บรรจุแอปพลิเคชันและโมเดลให้ deploy ได้ง่าย",
        ],
        y=4.15,
        size=18,
    )
    add_footer(slide, 4)

    # 5
    slide = prs.slides.add_slide(blank)
    add_title(slide, "ผลการ Optimization")
    add_table(
        slide,
        [
            ["Runtime", "Model Size", "Avg Latency", "P95 Latency"],
            ["PyTorch", "21.44 MB", "34.43 ms", "37.51 ms"],
            ["ONNX", "21.49 MB", "55.58 ms", "63.64 ms"],
            ["ONNX Quantized", "11.56 MB", "47.00 ms", "56.47 ms"],
        ],
        0.9,
        1.55,
        11.5,
        1.65,
    )
    add_card(slide, 1.0, 4.05, 5.2, 1.35, "ผลลัพธ์สำคัญ", "ONNX Quantized ลดขนาดโมเดลลงประมาณ 46% เมื่อเทียบกับ ONNX ปกติ", GREEN)
    add_card(slide, 7.0, 4.05, 5.2, 1.35, "ข้อสังเกต", "Latency ดีขึ้นกว่า ONNX ปกติ แต่ PyTorch บนเครื่องทดสอบยังเร็วกว่า", ORANGE)
    add_footer(slide, 5)

    # 6
    slide = prs.slides.add_slide(blank)
    add_title(slide, "API และ Error Handling")
    add_bullets(
        slide,
        [
            "Endpoint หลัก: /predict รับไฟล์รูปภาพแบบ multipart/form-data",
            "รองรับไฟล์ JPEG, PNG และ WebP",
            "ตรวจสอบไฟล์ว่าง ไฟล์เสีย และขนาดไฟล์ก่อนเข้าโมเดล",
            "ตอบกลับ 400 Bad Request สำหรับ input ที่ไม่ถูกต้อง",
            "ตอบกลับ 413 Request Entity Too Large เมื่อไฟล์มีขนาดใหญ่เกินกำหนด",
        ],
        size=19,
    )
    add_footer(slide, 6)

    # 7
    slide = prs.slides.add_slide(blank)
    add_title(slide, "ผลการทดสอบ API")
    add_card(slide, 0.9, 1.55, 3.5, 1.3, "Unit Test", "ผ่าน 3/3 tests", GREEN)
    add_card(slide, 4.9, 1.55, 3.5, 1.3, "Docker Test", "API รันและทำนายผลได้จริง", GREEN)
    add_card(slide, 8.9, 1.55, 3.5, 1.3, "Prediction", "runtime: onnx-quantized", GREEN)
    add_table(
        slide,
        [
            ["Input", "ผลลัพธ์"],
            ["sample_images/benchmark.png", "tabby cat, Egyptian cat, tiger cat"],
            ["README.md เป็น text/plain", "400 Bad Request"],
        ],
        1.2,
        3.6,
        10.8,
        1.25,
    )
    add_footer(slide, 7)

    # 8
    slide = prs.slides.add_slide(blank)
    add_title(slide, "ผลการทดสอบโหลด")
    add_table(
        slide,
        [
            ["หัวข้อ", "ผลการทดสอบ"],
            ["จำนวน request", "30"],
            ["Concurrent users", "10"],
            ["Success", "30/30"],
            ["Throughput", "10.32 requests/second"],
            ["Average Latency", "874.90 ms"],
            ["P95 Latency", "1187.96 ms"],
        ],
        1.3,
        1.4,
        10.5,
        3.0,
    )
    add_bullets(
        slide,
        [
            "เมื่อจำนวนผู้ใช้พร้อมกันเพิ่มขึ้น latency จะสูงขึ้นตามภาระ CPU",
            "หากเพิ่ม threads มากเกินไปอาจพบ Socket closed ซึ่งเป็นจุดที่ระบบเริ่มรับโหลดไม่ไหว",
        ],
        y=5.0,
        size=17,
    )
    add_footer(slide, 8)

    # 9
    slide = prs.slides.add_slide(blank)
    add_title(slide, "CI/CD Pipeline")
    add_card(slide, 1.0, 1.8, 2.3, 1.2, "Push Code", "ส่งโค้ดขึ้น GitHub")
    add_card(slide, 3.8, 1.8, 2.3, 1.2, "GitHub Actions", "ติดตั้ง dependencies")
    add_card(slide, 6.6, 1.8, 2.3, 1.2, "Unit Test", "รัน pytest")
    add_card(slide, 9.4, 1.8, 2.3, 1.2, "Deploy", "Hugging Face Spaces")
    add_bullets(
        slide,
        [
            "Workflow อยู่ที่ .github/workflows/ci-cd.yml",
            "รัน unit test ทุกครั้งที่ push หรือ pull request ไปยัง main",
            "รองรับการ deploy อัตโนมัติเมื่อกำหนด HF_TOKEN และ HF_SPACE_REPO_ID",
        ],
        y=4.15,
        size=18,
    )
    add_footer(slide, 9)

    # 10
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Live Demo และสรุปผล")
    add_bullets(
        slide,
        [
            "Demo: เปิด API บน Docker และเรียก /predict ด้วยไฟล์รูปภาพจริง",
            "ผลลัพธ์ที่คาดหวัง: API ตอบ JSON พร้อม model, runtime และ predictions",
            "ระบบใช้โมเดล apple/mobilevit-small เวอร์ชัน ONNX Quantized",
            "โปรเจกต์มี deliverables ครบ: Source Code, README, CI/CD, Docker, Postman, JMeter และ Report PDF",
        ],
        size=19,
    )
    add_footer(slide, 10)

    prs.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
